"""Excel 导出单元测试。"""
from __future__ import annotations

from datetime import date
from io import BytesIO

from openpyxl import load_workbook

from core import exporter, worktime


def _load(data: bytes):
    return load_workbook(BytesIO(data))


def test_export_schedule_content():
    rows = [
        {"work_date": date(2026, 3, 10), "weekday": "周二", "employee_name": "张三",
         "position": "前厅", "start_time": "09:00", "end_time": "17:00",
         "break_minutes": 30, "effective_hours": 7.5, "note": ""},
        {"work_date": date(2026, 3, 10), "weekday": "周二", "employee_name": "李四",
         "position": "后厨", "start_time": "19:00", "end_time": "01:00",
         "break_minutes": 0, "effective_hours": 6.0, "note": "高峰支援"},
    ]
    data = exporter.export_schedule(rows, "测试门店", date(2026, 3, 9), date(2026, 3, 15))
    ws = _load(data).active
    assert ws.cell(1, 1).value == "测试门店 排班表（2026-03-09 至 2026-03-15）"
    assert [ws.cell(2, i).value for i in range(1, 9)] == [
        "日期", "星期", "员工", "岗位", "班次时段", "休息(分钟)", "有效工时(小时)", "备注"]
    assert ws.cell(3, 3).value == "张三"
    assert ws.cell(3, 7).value == 7.5
    assert ws.cell(4, 5).value == "19:00–01:00"  # 跨日班次原样输出
    assert ws.max_row == 4  # 标题 + 表头 + 2 条记录


def test_export_hours_totals():
    from datetime import timedelta
    summary = worktime.summarize(
        [{"employee_id": 1, "employee_name": "张三", "position": "前厅",
          "employment_type": "fulltime", "work_date": date(2026, 3, 10),
          "start_time": "09:00", "end_time": "17:00", "break_minutes": 30},
         {"employee_id": 1, "employee_name": "张三", "position": "前厅",
          "employment_type": "fulltime", "work_date": date(2026, 3, 11),
          "start_time": "09:00", "end_time": "17:00", "break_minutes": 30}],
        date(2026, 3, 9), date(2026, 3, 15))
    day_columns = [((date(2026, 3, 9) + timedelta(days=i)).isoformat(), f"周{i+1}")
                   for i in range(7)]
    data = exporter.export_hours("测试门店", "2026-03-09 至 2026-03-15",
                                 summary, day_columns)
    ws = _load(data).active
    assert ws.cell(3, 4).value == 2              # 班次数（第 4 列）
    assert ws.cell(3, 7).value == 15.0           # 有效工时 16h − 1h 休息
    assert ws.cell(3, 9).value == 7.5            # 周二列（第 8 列起为逐日明细）
    total_row = ws.max_row
    assert ws.cell(total_row, 1).value == "合计"
    assert ws.cell(total_row, 7).value == 15.0


def test_export_empty_summary_still_valid():
    data = exporter.export_hours("测试门店", "空周期", [])
    ws = _load(data).active
    assert ws.cell(1, 1).value == "测试门店 工时统计（空周期）"
