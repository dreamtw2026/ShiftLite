"""Excel 报表导出（openpyxl）：排班表与工时统计表。

配色沿用系统 Material 3 米 beige 主题，表头使用主色容器色。
"""
from __future__ import annotations

from datetime import date as Date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# 与前端 theme.css 一致的主题色
THEME = {
    "title": "78664F",        # 主色
    "header_fill": "E9DDC8",  # 主色容器
    "surface": "FFFDF8",      # 表面
    "text": "292720",         # 主要文字
    "subtext": "68645C",      # 次要文字
    "border": "DED7CA",       # 边框
    "total_fill": "F1EBDD",   # 合计行底色
}

_THIN = Side(style="thin", color=THEME["border"])
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)


def _apply_widths(ws, widths: list[float]) -> None:
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width


def _title_row(ws, text: str, col_span: int) -> int:
    """写入标题行，返回下一可用行号。"""
    ws.cell(row=1, column=1, value=text)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=col_span)
    cell = ws.cell(row=1, column=1)
    cell.font = Font(bold=True, size=14, color=THEME["title"])
    cell.alignment = _LEFT
    ws.row_dimensions[1].height = 26
    return 2


def _header_row(ws, row: int, headers: list[str]) -> None:
    for col, name in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=name)
        cell.font = Font(bold=True, size=11, color=THEME["text"])
        cell.fill = PatternFill("solid", fgColor=THEME["header_fill"])
        cell.alignment = _CENTER
        cell.border = _BORDER
    ws.row_dimensions[row].height = 22


def export_schedule(rows: list[dict], store_name: str,
                    start: Date, end: Date) -> bytes:
    """导出排班表。

    rows 每条含：work_date(Date)、weekday、employee_name、position、
    start_time、end_time、break_minutes、effective_hours(float)、note。
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "排班表"
    headers = ["日期", "星期", "员工", "岗位", "班次时段",
               "休息(分钟)", "有效工时(小时)", "备注"]
    next_row = _title_row(ws, f"{store_name} 排班表（{start.isoformat()} 至 {end.isoformat()}）",
                          len(headers))
    _header_row(ws, next_row, headers)
    r = next_row + 1
    for row in rows:
        values = [
            row["work_date"].isoformat(),
            row["weekday"],
            row["employee_name"],
            row["position"] or "",
            f"{row['start_time']}–{row['end_time']}",
            int(row["break_minutes"] or 0),
            round(float(row["effective_hours"]), 2),
            row.get("note") or "",
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=col, value=value)
            cell.border = _BORDER
            cell.font = Font(size=10, color=THEME["text"])
            cell.alignment = _LEFT if col in (3, 4, 8) else _CENTER
        r += 1
    _apply_widths(ws, [12, 7, 10, 12, 16, 11, 14, 20])
    ws.freeze_panes = ws.cell(row=next_row + 1, column=1)
    ws.sheet_view.showGridLines = False

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def export_hours(store_name: str, period_label: str,
                 summary: list[dict], day_columns: list[tuple[str, str]] | None = None) -> bytes:
    """导出工时统计表。

    summary 每条为 worktime.summarize 的输出项。
    day_columns 为 [(ISO日期, 显示名), ...]，周视图传入周一至周日。
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "工时统计"
    headers = ["员工", "岗位", "类型", "班次数", "排班时长(小时)",
               "休息时长(小时)", "有效工时(小时)"]
    if day_columns:
        headers += [label for _, label in day_columns]
    next_row = _title_row(ws, f"{store_name} 工时统计（{period_label}）", len(headers))
    _header_row(ws, next_row, headers)
    r = next_row + 1
    for item in summary:
        values = [
            item["name"],
            item["position"] or "",
            item["type_label"],
            item["shifts"],
            round(item["scheduled_minutes"] / 60, 2),
            round(item["break_minutes"] / 60, 2),
            round(item["effective_minutes"] / 60, 2),
        ]
        if day_columns:
            for day_iso, _ in day_columns:
                values.append(round(item["by_date"].get(day_iso, 0) / 60, 2))
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=col, value=value)
            cell.border = _BORDER
            cell.font = Font(size=10, color=THEME["text"])
            cell.alignment = _LEFT if col in (1, 2) else _CENTER
        r += 1
    # 合计行
    total = ["合计", "", "", sum(i["shifts"] for i in summary),
             round(sum(i["scheduled_minutes"] for i in summary) / 60, 2),
             round(sum(i["break_minutes"] for i in summary) / 60, 2),
             round(sum(i["effective_minutes"] for i in summary) / 60, 2)]
    if day_columns:
        total += [round(sum(i["by_date"].get(d, 0) for i in summary) / 60, 2)
                  for d, _ in day_columns]
    for col, value in enumerate(total, start=1):
        cell = ws.cell(row=r, column=col, value=value)
        cell.border = _BORDER
        cell.font = Font(bold=True, size=10, color=THEME["text"])
        cell.fill = PatternFill("solid", fgColor=THEME["total_fill"])
        cell.alignment = _LEFT if col in (1, 2) else _CENTER
    widths = [10, 12, 7, 8, 14, 13, 14] + [9] * len(day_columns or [])
    _apply_widths(ws, widths)
    ws.freeze_panes = ws.cell(row=next_row + 1, column=1)
    ws.sheet_view.showGridLines = False

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
