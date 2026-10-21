from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

out = Path('docs/团队信息表_待填写.xlsx')
wb = Workbook()
ws = wb.active
ws.title = '项目基本信息'
ws.append(['字段', '填写内容', '填写说明'])
rows = [
('项目名称', '轻排班（ShiftLite）', '按正式报名名称确认'),
('团队名称', None, '与项目书封面保持一致'),
('比赛名称', None, '填写主办方公布的完整名称'),
('参赛赛道及组别', None, '按正式报名要求填写'),
('所属学校', None, '学校全称'),
('学院或院系', None, '填写所属院系'),
('项目负责人', None, '与成员信息中的负责人一致'),
('联系人电话', None, '文本格式，保留前导零'),
('联系人邮箱', None, '常用邮箱'),
('项目版本', None, '填写本次提交版本'),
('填表日期', None, 'YYYY-MM-DD'),
('项目简介', None, '简述服务对象、解决的问题及主要功能'),
('作品或演示地址', None, '按主办方要求填写，无则留空'),
('备注', None, '其他需要说明的信息'),
]
for row in rows: ws.append(row)
ws.column_dimensions['A'].width = 24
ws.column_dimensions['B'].width = 54
ws.column_dimensions['C'].width = 52
ws = wb.create_sheet('团队成员')
ws.append(['序号','姓名','团队角色','学校','学院或专业','年级或班级','学号（如要求）','联系电话','邮箱','具体分工','负责成果或模块'])
for i in range(1, 9): ws.append([i] + [None] * 10)
dv = DataValidation(type='list', formula1='"负责人,成员"', allow_blank=True)
ws.add_data_validation(dv)
dv.add('C2:C9')
ws = wb.create_sheet('指导教师')
ws.append(['序号','姓名','学校或单位','学院或部门','职称或职务','联系电话','邮箱','指导内容'])
for i in range(1, 4): ws.append([i] + [None] * 7)
ws = wb.create_sheet('填写说明')
ws.append(['事项','说明'])
for row in [
('模板性质','通用草稿，不代表主办方官方表格；如有官方模板，以官方模板为准。'),
('空白字段','未核实信息均留空，按实际情况填写；不适用的字段可写“不适用”。'),
('成员人数','预留8行成员、3行教师，仅为填写空间，不代表参赛人数要求；可按规定增删行。'),
('信息一致','项目名称、团队名称、成员顺序、负责人及指导教师应与项目书和报名系统一致。'),
('分工填写','填写具体负责模块、工作内容或成果；不要只写“参与开发”。'),
('个人信息','学号等字段仅在主办方要求时填写；对外公开材料按需隐去联系方式。'),
('提交前核对','核对比赛、赛道、组别、人数限制、签字盖章要求及文件命名规则。'),
]: ws.append(row)
ws.column_dimensions['A'].width = 22
ws.column_dimensions['B'].width = 100
for ws in wb:
    ws.freeze_panes = 'A2'
    for c in ws[1]:
        c.font = Font(name='Microsoft YaHei', bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='526B55')
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name='Microsoft YaHei', size=11)
            c.alignment = Alignment(vertical='top', wrap_text=True)
            if c.column > 1: c.number_format = '@'
    for row in range(2, ws.max_row + 1): ws.row_dimensions[row].height = 36
    if ws.title in ['团队成员', '指导教师']:
        ws.auto_filter.ref = ws.dimensions
        for col in ws.columns: ws.column_dimensions[col[0].column_letter].width = 22
        ws.column_dimensions['A'].width = 8
wb.save(out)
check = load_workbook(out)
assert check.sheetnames == ['项目基本信息','团队成员','指导教师','填写说明']
assert check['团队成员']['B2'].value is None
assert check['项目基本信息']['B2'].value == '轻排班（ShiftLite）'
print(str(out))
